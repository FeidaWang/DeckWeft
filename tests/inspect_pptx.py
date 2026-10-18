"""Optional independent PPTX inspection. Requires python-pptx and lxml.
Run after npm run demo. Does not claim Microsoft PowerPoint UI acceptance.
"""
from pathlib import Path
from zipfile import ZipFile
import json,posixpath
from lxml import etree
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1]
results=[]
for file in sorted((ROOT/'dist/demo').glob('*.pptx')):
    presentation=Presentation(file)
    native_text=sum(1 for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame and shape.text.strip())
    picture_count=sum(1 for slide in presentation.slides for shape in slide.shapes if shape.shape_type==13)
    with ZipFile(file) as archive:
        assert archive.testzip() is None
        parts=set(archive.namelist())
        for name in parts:
            if name.endswith(('.xml','.rels')):etree.fromstring(archive.read(name))
            if name.endswith('.rels'):
                base=posixpath.dirname(posixpath.dirname(name))
                for rel in etree.fromstring(archive.read(name)):
                    if rel.get('TargetMode')!='External':
                        target=posixpath.normpath(posixpath.join(base,rel.get('Target')))
                        assert target in parts,(name,target)
        assert native_text>0 and picture_count==0
    results.append({'file':file.name,'slides':len(presentation.slides),'nativeTextShapes':native_text,'pictureShapes':picture_count,'zipCRC':'passed','xmlParsing':'passed','relationshipTargets':'passed'})
assert len(results)==2,'Generate the two language editions with npm run demo first.'
(ROOT/'docs/pptx-evidence.json').write_text(json.dumps({'status':'passed','scope':'independent package/parser inspection, not Microsoft PowerPoint UI validation','files':results},indent=2)+'\n')
print(json.dumps(results,indent=2))
